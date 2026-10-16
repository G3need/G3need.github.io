#!/usr/bin/env python3
"""
ZED Dossier Builder — UK lane.
Reads the UK dossier docx, fixes placeholders with Ahmed's real data,
reframes Section 2 as employer pitch (pre-CoS), outputs corrected .docx.
"""
import os
from docx import Document

SRC = r"C:\Users\Geneed\AppData\Local\hermes\cache\documents\doc_d7f5a5760aa6_ZED_Dossier_Pilot_-_United_Kingdom_Tech___Scaleup_Skilled_Worker__Ahmed_Geneed_.docx"
OUT_DIR = r"C:\Users\Geneed\Desktop\draconis-hub\zed"

# Real data (verified: business WA +201551961964; personal email used on CV/LinkedIn)
REAL = {
    "email": "ahmedgeneed@gmail.com",
    "phone": "+20 155 196 1964",
    "location": "Ismailia / Cairo, Egypt",
}
# GitHub: only real repos. Ahmed's GitHub org/user = G3need (github.com/G3need).
# The dossier's links (agentic-orchestrator, dsp-signal-suite) are not verified repos — replace with flagship real ones.
GITHUB_LINE = ("Selected Projects: HORUS (agentic AI systems, github.com/G3need) · "
               "Sovereign-DSP (real-time DSP suite) · FOOTY4ALL (edge computer-vision analytics). "
               "Full portfolio: github.com/G3need")

REPLACEMENTS = [
    ("Location: Cairo, Egypt | Email: ahmed.geneed@example.com | Phone: +20 100 000 000",
     f"Location: {REAL['location']} | Email: {REAL['email']} | Phone: {REAL['phone']}"),
    ("Sponsoring Entity: Person (Sponsor Name)", "Sponsoring Entity: [To be secured — see sponsor target list, zed/uk_sponsor_targets.md]"),
    ("Sponsor Licence Number: File", "Sponsor Licence Number: [pending sponsor]"),
    ("Certificate of Sponsorship (CoS) Reference: File", "Certificate of Sponsorship (CoS) Reference: [pending sponsor]"),
]

def main():
    doc = Document(SRC)

    # 1) apply paragraph replacements
    changed = 0
    for p in doc.paragraphs:
        for old, new in REPLACEMENTS:
            if old in p.text:
                # replace text while keeping runs' style of first run
                for run in p.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        changed += 1
                # fallback if split across runs
                if old in p.text:
                    p.text = p.text.replace(old, new)
                    changed += 1

    # 2) GitHub line fix
    for p in doc.paragraphs:
        if "agentic-orchestrator" in p.text or "dsp-signal-suite" in p.text:
            p.text = GITHUB_LINE
            changed += 1

    # 3) reframe Section 2 header: it's a pre-CoS pitch, not a UKVI support letter
    for p in doc.paragraphs:
        if "OFFICIAL SPONSOR SUPPORT & COVER LETTER" in p.text:
            p.text = "EMPLOYER PITCH & PRE-COS COVER LETTER (send to sponsor targets BEFORE Certificate of Sponsorship exists)"
            changed += 1
        if p.text.strip().startswith("I am writing to formally submit and support the Skilled Worker Visa application"):
            p.text = ("I am writing to introduce Mr. Ahmed Geneed as a candidate for an AI & Software Architect role under "
                      "SOC Code 2134 (Programmers and Software Development Professionals). This letter supports a future "
                      "Skilled Worker application once a Certificate of Sponsorship is issued.")
            changed += 1
        if p.text.strip().startswith("Mr. Geneed’s background, formal credentials"):
            p.text = ("Mr. Geneed’s background, formal credentials, and specialized engineering skills position him to deliver "
                      "immediate value to a UK technology team. We welcome the opportunity to discuss sponsorship under the "
                      "Skilled Worker route.")
            changed += 1

    # 4) closing signature fix
    for p in doc.paragraphs:
        if p.text.strip() == "Person":
            p.text = "Ahmed Geneed"
            changed += 1
        elif p.text.strip() == "Place":
            p.text = REAL["location"]
            changed += 1
        elif p.text.strip() == "Date":
            import datetime
            p.text = datetime.date.today().isoformat()
            changed += 1

    out = os.path.join(OUT_DIR, "ZED_UK_Dossier_Ahmed_Geneed_CORRECTED.docx")
    doc.save(out)
    print(f"changed {changed} elements -> {out}")

    # verify
    doc2 = Document(out)
    bad = 0
    for p in doc2.paragraphs:
        t = p.text
        if "example.com" in t or "+20 100 000 000" in t or "Sponsor Name)" in t:
            bad += 1
            print("STILL BAD:", t[:100])
    print("verification:", "PASS" if bad == 0 else f"FAIL ({bad} placeholders remain)")

if __name__ == "__main__":
    main()
