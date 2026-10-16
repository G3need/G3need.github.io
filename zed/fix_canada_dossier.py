#!/usr/bin/env python3
"""
ZED Dossier Builder — Canada lane.
Fixes Canada dossier placeholders/typos with Ahmed's real data.
"""
import os, datetime
from docx import Document

SRC = r"C:\Users\Geneed\AppData\Local\hermes\cache\documents\doc_888e06b2b97a_ZED_Dossier_Pilot_-_Canada_Tech___STEM_Express_Entry__Ahmed_Geneed_.docx"
OUT_DIR = r"C:\Users\Geneed\Desktop\draconis-hub\zed"

REPLACEMENTS = [
    ("ahmed.geneed@email.com", "ahmedgeneed@gmail.com"),
    ("Engineerd ", "Engineered "),
    # IRCC explicitly says do not include a cover letter with e-APR — reframe header to LOE-only usage
    ("Official Statement of Purpose (SOP) & Letter of Explanation (LOE)",
     "Letter of Explanation (LOE) — IRCC-compliant (no separate SOP is submitted with e-APR)"),
]

def main():
    doc = Document(SRC)
    changed = 0
    for p in doc.paragraphs:
        for old, new in REPLACEMENTS:
            if old in p.text:
                if p.runs and old in p.runs[0].text:
                    for run in p.runs:
                        run.text = run.text.replace(old, new)
                else:
                    p.text = p.text.replace(old, new)
                changed += 1
    # signature "Place/Date" pattern if present
    for p in doc.paragraphs:
        if p.text.strip() == "Place":
            p.text = "Ismailia, Egypt"
            changed += 1
        elif p.text.strip() == "Date":
            p.text = datetime.date.today().isoformat()
            changed += 1
    out = os.path.join(OUT_DIR, "ZED_Canada_Dossier_Ahmed_Geneed_CORRECTED.docx")
    doc.save(out)
    print(f"changed {changed} -> {out}")
    # verify no known placeholders remain
    doc2 = Document(out)
    bad = [p.text[:80] for p in doc2.paragraphs if "email.com" in p.text or "Engineerd " in p.text]
    print("verification:", "PASS" if not bad else f"FAIL: {bad}")

if __name__ == "__main__":
    main()
